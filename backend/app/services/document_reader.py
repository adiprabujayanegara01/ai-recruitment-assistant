"""Validasi file & ekstraksi teks dari PDF / DOCX."""
import io
import re
import zipfile

from app.core.config import get_settings

ALLOWED = {".pdf": "pdf", ".docx": "docx"}


class FileValidationError(ValueError):
    pass


def validate_upload(filename: str, data: bytes) -> str:
    """Cek ekstensi, ukuran, dan magic bytes (bukan hanya ekstensi). Mengembalikan tipe file."""
    settings = get_settings()
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in ALLOWED:
        raise FileValidationError("Format tidak didukung. Gunakan PDF atau DOCX.")
    if len(data) == 0:
        raise FileValidationError("File kosong.")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise FileValidationError(f"Ukuran file melebihi {settings.max_upload_mb} MB.")
    kind = ALLOWED[ext]
    if kind == "pdf" and not data.startswith(b"%PDF-"):
        raise FileValidationError("Isi file bukan PDF yang valid.")
    if kind == "docx":
        if not data.startswith(b"PK"):
            raise FileValidationError("Isi file bukan DOCX yang valid.")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise FileValidationError("Isi file bukan DOCX yang valid.")
        except zipfile.BadZipFile as e:
            raise FileValidationError("Isi file bukan DOCX yang valid.") from e
    return kind


def _clean(text: str) -> str:
    text = text.replace("\u00a0", " ").replace("\u2022", "•").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_text(data: bytes, kind: str) -> str:
    if kind == "pdf":
        import fitz  # PyMuPDF

        with fitz.open(stream=data, filetype="pdf") as doc:
            text = "\n".join(page.get_text("text") for page in doc)
    else:
        from docx import Document

        d = Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:  # CV sering memakai tabel
            for row in table.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        text = "\n".join(parts)
    text = _clean(text)
    if len(text) < 30:
        raise FileValidationError(
            "Teks tidak dapat diekstrak (mungkin PDF hasil scan). Gunakan file dengan teks yang dapat dipilih."
        )
    return text


def extract_text_from_plain(data: bytes) -> str:
    return _clean(data.decode("utf-8", errors="ignore"))
